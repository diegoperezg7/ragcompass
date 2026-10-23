"""Counts checked by hand against the content written in examples/generate.py."""

import pytest

from ragcompass.ingest import extract
from ragcompass.layout import Layout
from ragcompass.models import BlockType, DocKind
from ragcompass.profile import build_profile
from tests.conftest import CORPUS


def profile(name, tokenizer):
    path = CORPUS / name
    extracted = extract(path)
    assert extracted.error is None
    return build_profile(name, path.stat().st_size, extracted, Layout(extracted.blocks, tokenizer))


def test_text_pdf(tokenizer):
    p = profile("hx200-maintenance-manual.pdf", tokenizer)
    assert p.kind == DocKind.PDF_TEXT
    assert p.headings == 8  # 1 title, 5 numbered sections, 2 subsections
    assert p.tables == 2
    assert p.table_rows == 5 + 7  # maintenance tasks + verification points
    assert p.list_items == 4 + 4  # safety rules + troubleshooting codes
    assert p.images == 1  # wiring diagram
    assert p.scanned_pages == 0
    assert p.language == "en"


def test_scanned_pdf_has_no_text_layer(tokenizer):
    p = profile("scanned-invoice.pdf", tokenizer)
    assert p.kind == DocKind.PDF_SCANNED
    assert p.pages == p.scanned_pages == 1
    assert p.tokens == 0
    assert p.images == 0  # the scan itself is not counted as a figure


def test_docx(tokenizer):
    p = profile("politica-devoluciones.docx", tokenizer)
    assert p.kind == DocKind.DOCX
    assert p.headings == 5
    assert p.list_items == 4
    assert (p.tables, p.table_rows) == (1, 4)
    assert p.images == 1
    assert p.language == "es"


def test_markdown(tokenizer):
    p = profile("onboarding-handbook.md", tokenizer)
    assert p.headings == 6  # 1 title, 4 sections, 1 subsection
    assert (p.tables, p.table_rows) == (1, 3)
    assert p.list_items == 4 + 3
    assert p.code_blocks == 1


def test_html_strips_navigation(tokenizer):
    path = CORPUS / "hx200-faq.html"
    extracted = extract(path)
    text = " ".join(b.text for b in extracted.blocks)
    assert "Products" not in text  # <nav>
    assert "fictional" not in text  # <footer>
    assert "analytics" not in text  # <script>
    p = profile("hx200-faq.html", tokenizer)
    assert p.headings == 6
    assert p.images == 2


def test_csv(tokenizer):
    p = profile("product-catalog.csv", tokenizer)
    assert p.kind == DocKind.CSV
    assert (p.tables, p.table_rows) == (1, 40)
    assert p.identifiers.dates == 40  # one last_updated date per row


def test_xlsx_one_table_per_sheet(tokenizer):
    p = profile("warehouse-stock.xlsx", tokenizer)
    assert (p.headings, p.tables, p.table_rows) == (2, 2, 8 + 6)


@pytest.mark.parametrize(
    ("name", "kind"),
    [("ticket_export.py", DocKind.CODE), ("wiring-diagram.png", DocKind.IMAGE)],
)
def test_other_kinds(name, kind):
    extracted = extract(CORPUS / name)
    assert extracted.kind == kind
    if kind == DocKind.CODE:
        assert extracted.blocks[0].type == BlockType.CODE
    else:
        assert extracted.images == 1


def test_broken_file_is_reported_not_raised(tmp_path):
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"not a pdf")
    assert extract(broken).error


def test_text_encodings(tmp_path):
    from ragcompass.ingest import read_text

    text = "Café, déjà vu, 5 €"
    cases = {
        "utf8.txt": text.encode("utf-8"),
        "bom.txt": b"\xef\xbb\xbf" + text.encode("utf-8"),
        "utf16.txt": text.encode("utf-16"),  # Python writes the BOM
        "cp1252.txt": text.encode("cp1252"),
    }
    for name, data in cases.items():
        (tmp_path / name).write_bytes(data)
        assert read_text(tmp_path / name) == text, name


def test_data_files_are_text_not_code(tmp_path):
    (tmp_path / "settings.yaml").write_text("# comment\nkey: value\n\n- item\n")
    extracted = extract(tmp_path / "settings.yaml")
    assert extracted.kind == DocKind.TEXT
    assert {b.type for b in extracted.blocks} == {BlockType.PARAGRAPH}  # no headings or lists


def test_own_outputs_are_skipped(tmp_path):
    from ragcompass.analyze import analyze
    from ragcompass.ingest import iter_files
    from ragcompass.report import to_json, to_markdown, to_yaml
    from ragcompass.tokens import RegexTokenizer

    (tmp_path / "notes.md").write_text("# Notes\n\nSome text about the calibrator.\n")
    report = analyze(tmp_path, tokenizer=RegexTokenizer())
    (tmp_path / "any-name.yaml").write_text(to_yaml(report))
    (tmp_path / "report.md").write_text(to_markdown(report))
    (tmp_path / "report.json").write_text(to_json(report))
    assert [p.name for p in iter_files(tmp_path)] == ["notes.md"]


def test_table_cell_budget(tmp_path, monkeypatch, tokenizer):
    from openpyxl import Workbook

    from ragcompass.ingest import tabular

    workbook = Workbook()
    for i in range(50):
        workbook.active.append([f"r{i}", i, i * 2])
    workbook.save(tmp_path / "big.xlsx")
    monkeypatch.setattr(tabular, "MAX_TABLE_CELLS", 30)  # 10 rows of 3 cells
    extracted = extract(tmp_path / "big.xlsx")
    assert extracted.blocks[1].rows is not None and len(extracted.blocks[1].rows) == 10
    assert extracted.warnings == ["truncated after 30 table cells"]


def _scan_pdf(path, pages):
    """Build a PDF from page specs: lists of ("image", rect fractions) or ("text", str, y)."""
    import pymupdf

    document = pymupdf.open()
    for spec in pages:
        page = document.new_page(width=595, height=842)
        for item in spec:
            if item[0] == "image":
                x0, y0, x1, y1 = item[1]
                pix = pymupdf.Pixmap(pymupdf.csGRAY, pymupdf.IRect(0, 0, 60, 80), False)
                pix.clear_with(200)
                page.insert_image(
                    pymupdf.Rect(x0 * 595, y0 * 842, x1 * 595, y1 * 842), pixmap=pix, keep_proportion=False
                )
            else:
                page.insert_text((40, item[2]), item[1], fontsize=9)
    document.save(path)


FULL = ("image", (0, 0, 1, 1))


def test_blank_page_does_not_turn_a_scan_into_mixed(tmp_path):
    _scan_pdf(tmp_path / "scan.pdf", [[FULL], []])
    extracted = extract(tmp_path / "scan.pdf")
    assert (extracted.kind, extracted.scanned_pages) == (DocKind.PDF_SCANNED, 1)


def test_stamp_on_a_scan_still_needs_ocr(tmp_path):
    _scan_pdf(tmp_path / "scan.pdf", [[FULL, ("text", "APPROVED 2025-11-04 QA-DEPT / page 1 of 1", 820)]])
    assert extract(tmp_path / "scan.pdf").kind == DocKind.PDF_SCANNED


def test_scan_split_into_strips(tmp_path):
    strips = [("image", (0, 0, 1, 0.34)), ("image", (0, 0.34, 1, 0.67)), ("image", (0, 0.67, 1, 1))]
    _scan_pdf(tmp_path / "scan.pdf", [strips])
    assert extract(tmp_path / "scan.pdf").kind == DocKind.PDF_SCANNED


def test_text_page_with_large_figure_is_not_a_scan(tmp_path):
    lines = [
        ("text", "Body text of a manual page with enough words to be a real paragraph.", 600 + 12 * i)
        for i in range(18)
    ]
    _scan_pdf(tmp_path / "doc.pdf", [[("image", (0, 0, 1, 0.6)), *lines]])
    extracted = extract(tmp_path / "doc.pdf")
    assert (extracted.kind, extracted.images) == (DocKind.PDF_TEXT, 1)
