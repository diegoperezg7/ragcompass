"""Build the sample corpus in examples/corpus.

Everything here is synthetic and written for this repository (CC0). The company,
products and people are fictional. Run with: uv run python examples/generate.py
"""

from __future__ import annotations

import csv
import tempfile
from pathlib import Path

import docx
import pymupdf
from docx.shared import Cm
from openpyxl import Workbook

OUT = Path(__file__).parent / "corpus"

CSS = """
body { font-family: sans-serif; font-size: 10pt; line-height: 1.35; }
h1 { font-size: 20pt; margin-bottom: 6pt; }
h2 { font-size: 14pt; margin-top: 12pt; margin-bottom: 4pt; }
h3 { font-size: 12pt; margin-top: 8pt; }
p { margin-bottom: 6pt; }
table { border-collapse: collapse; margin: 6pt 0; }
th, td { border: 0.8pt solid #333; padding: 3pt 6pt; font-size: 9pt; }
th { background-color: #e6e6ef; }
li { margin-bottom: 2pt; }
"""

MANUAL = """
<h1>HX-200 Bench Calibrator: Maintenance Manual</h1>
<p>This manual covers routine maintenance of the HX-200 bench calibrator built by Halden Instruments.
It is written for service technicians who already hold a level 2 electrical safety certificate.
Read the safety section before opening the enclosure. Revision R4 replaces R3 from 2024-03-18.</p>

<h2>1. Safety</h2>
<p>The calibrator stores energy in its output capacitors for up to ninety seconds after the mains
switch is turned off. Always wait two minutes, then confirm with a meter that the voltage across
terminals T1 and T2 is below 5 V before touching any internal part. Never bypass the interlock on
the rear panel: it is the only protection against the 1,000 V reference output.</p>
<ul>
<li>Wear insulated gloves rated for 1,000 V when the enclosure is open.</li>
<li>Disconnect the instrument from the mains and from any device under test.</li>
<li>Keep the service log within reach and record every intervention.</li>
<li>Report any burnt smell, discoloured board or swollen capacitor to the service lead.</li>
</ul>

<h2>2. Scheduled maintenance</h2>
<p>Preventive maintenance follows the operating hours shown on the status screen. Intervals assume
a clean laboratory at 23 degrees and 45 percent relative humidity. Halve every interval when the
instrument works in a production line or in a dusty workshop, because dust on the reference board
shifts the readings long before it triggers an alarm.</p>
<table>
<tr><th>Task</th><th>Interval</th><th>Part number</th><th>Time</th></tr>
<tr><td>Clean air filter</td><td>500 h</td><td>HX-FLT-021</td><td>10 min</td></tr>
<tr><td>Replace fan assembly</td><td>8,000 h</td><td>HX-FAN-118</td><td>35 min</td></tr>
<tr><td>Replace backup battery</td><td>24 months</td><td>BAT-CR2032-HX</td><td>15 min</td></tr>
<tr><td>Verify reference output</td><td>12 months</td><td>CAL-KIT-07</td><td>2 h</td></tr>
<tr><td>Replace output relay</td><td>100,000 cycles</td><td>RLY-4471-B</td><td>40 min</td></tr>
</table>
<p>The time column is the hands-on time for a trained technician. It does not include the warm-up
period of forty-five minutes that the instrument needs before any verification step.</p>

<h3>2.1 Cleaning the air filter</h3>
<p>Remove the four screws of the rear grille with a T10 driver. Slide the filter HX-FLT-021 out of
its rail and blow it clean from the inside towards the outside with dry compressed air at no more
than 2 bar. Replace it when the foam is torn or when it stays grey after cleaning. Refit the grille
and tighten the screws to 0.6 N·m. Record the cleaning in the service log with the operating hours.</p>

<h3>2.2 Replacing the fan assembly</h3>
<p>A worn fan announces itself with a rising pitch and with the warning W-212 on the status screen.
Order the fan assembly HX-FAN-118, which includes the guard and the connector. Unplug the fan cable
from connector J14 on the power board, remove the two retaining clips and lift the old fan out. The
new fan must blow towards the rear of the case; an arrow on its frame shows the airflow.</p>

<h2>3. Verification of the reference output</h2>
<p>Verification compares the output of the calibrator with a reference standard whose uncertainty
is at least four times smaller. Use the calibration kit CAL-KIT-07 and a reference multimeter that
has been calibrated within the last twelve months. Verify the seven points of the table below, in
the listed order, after the forty-five minute warm-up.</p>
<table>
<tr><th>Point</th><th>Nominal</th><th>Tolerance</th><th>Range</th></tr>
<tr><td>P1</td><td>10 mV</td><td>±2 µV</td><td>100 mV</td></tr>
<tr><td>P2</td><td>100 mV</td><td>±8 µV</td><td>100 mV</td></tr>
<tr><td>P3</td><td>1 V</td><td>±25 µV</td><td>1 V</td></tr>
<tr><td>P4</td><td>10 V</td><td>±180 µV</td><td>10 V</td></tr>
<tr><td>P5</td><td>100 V</td><td>±2.1 mV</td><td>100 V</td></tr>
<tr><td>P6</td><td>500 V</td><td>±12 mV</td><td>1,000 V</td></tr>
<tr><td>P7</td><td>1,000 V</td><td>±26 mV</td><td>1,000 V</td></tr>
</table>
<p>If a point falls outside its tolerance, repeat it once after five minutes. If it fails again,
the instrument must be adjusted by an accredited laboratory; field adjustment voids the certificate
issued with the instrument. Attach the verification record to the service log and label the
instrument with the date of the next verification.</p>

<h2>4. Troubleshooting</h2>
<p>Most faults reported by customers come from three sources: a blocked air filter, a dead backup
battery or a worn output relay. The status screen shows a code for each of them. Start with the
code, follow the table in this section and only then open the enclosure.</p>
<ul>
<li>W-212 fan speed low: clean the filter, then replace HX-FAN-118 if the warning stays.</li>
<li>E-301 clock lost after power off: replace the backup battery BAT-CR2032-HX.</li>
<li>E-415 output relay contact resistance high: replace relay RLY-4471-B.</li>
<li>E-520 reference drift: verify the output; if P3 or P4 fail, send the unit for adjustment.</li>
</ul>
<p>When none of these codes is shown and the readings are still unstable, check the grounding of
the bench first. A floating ground can add noise of several microvolts on the lower ranges, which
looks like a fault in the instrument but disappears as soon as the bench is bonded to earth.</p>
<p><img src="wiring-diagram.png" width="380" height="199"/></p>
<p>Figure 1. Simplified wiring of the output stage, with relay K1 between the reference board and
terminals T1 and T2.</p>

<h2>5. Spare parts and support</h2>
<p>Spare parts are shipped from the central warehouse within two working days. Quote the serial
number of the instrument, which starts with HX2 followed by seven digits, when you order. Support
tickets are answered by the service desk from Monday to Friday. Urgent tickets about a safety issue
are answered within four hours.</p>
"""

INVOICE = """
<h1>INVOICE</h1>
<p>Halden Instruments GmbH, Industriestrasse 12, 8005 Zurich</p>
<p>Invoice number: INV-2025-00871<br/>Date: 2025-11-04<br/>Customer: Brightline Labs, account C-30417</p>
<table>
<tr><th>Item</th><th>Description</th><th>Qty</th><th>Unit price</th><th>Total</th></tr>
<tr><td>HX-FAN-118</td><td>Fan assembly HX-200</td><td>2</td><td>84.00</td><td>168.00</td></tr>
<tr><td>HX-FLT-021</td><td>Air filter, pack of 5</td><td>3</td><td>19.50</td><td>58.50</td></tr>
<tr><td>RLY-4471-B</td><td>Output relay</td><td>1</td><td>132.00</td><td>132.00</td></tr>
<tr><td>CAL-SVC-12</td><td>Annual verification service</td><td>1</td><td>420.00</td><td>420.00</td></tr>
</table>
<p>Subtotal CHF 778.50. VAT 8.1% CHF 63.06. Total due CHF 841.56.</p>
<p>Payment within 30 days to IBAN CH93 0076 2011 6238 5295 7. Please quote INV-2025-00871.</p>
"""

HANDBOOK = """# Service Desk Onboarding Handbook

Welcome to the service desk of Halden Instruments. This handbook explains how tickets flow, which
tools you will use and what we expect from you in your first month. Keep it open during your first
shifts; most questions new colleagues ask are answered here.

## How a ticket flows

Every request from a customer becomes a ticket. A ticket starts in the triage queue, where the
on-call agent reads it, sets a priority and assigns a product. Tickets about the HX-200 bench
calibrator go to the instruments team; tickets about invoices go to billing. A ticket is only closed
when the customer confirms that the problem is solved or after five working days without an answer.

1. Read the ticket and check the serial number against the asset register.
2. Set the priority using the table below.
3. Assign the product and the team.
4. Reply to the customer within the first-response target.

| Priority | Meaning | First response | Example |
|---|---|---|---|
| P1 | Safety issue or production stopped | 4 hours | Burnt smell from an HX-200 |
| P2 | Instrument unusable, workaround exists | 1 working day | Error E-415 on the output |
| P3 | Question or minor fault | 3 working days | How to export the service log |

## Tools

We work with three tools. The ticketing system holds the conversation with the customer. The asset
register tells you which instruments a customer owns, when they were last verified and which
contract covers them. The knowledge base contains procedures written by the instruments team; link
the relevant article in your reply instead of copying it, so customers always see the latest version.

### Writing a good reply

A good reply restates the problem in one sentence, gives the next step and says when the customer
will hear from us again. Avoid internal codes in replies unless the customer used them first. When
you ask for information, ask for everything you need in one message: the serial number, the error
code shown on the status screen and the firmware version.

- Use the customer's name and the ticket number in the subject.
- Quote part numbers exactly, for example HX-FAN-118, never "the fan".
- Never promise a delivery date that the warehouse has not confirmed.

## Escalation

Escalate a ticket to the service lead when a safety issue is reported, when a customer asks for a
refund above CHF 2,000, or when the same instrument comes back with the same fault within ninety
days. Escalation does not mean handing over the ticket: you stay the contact for the customer and
the lead helps you decide.

## Your first month

During your first week you will shadow a senior agent and answer P3 tickets only. In the second week
you will take P2 tickets with a reviewer. From the third week you join the on-call rotation for
triage. At the end of the month we review ten of your tickets together and agree on what to work on.

```bash
# Export your open tickets for the weekly review
desk export --assignee me --status open --format csv > my-tickets.csv
```
"""

FAQ = """<!doctype html>
<html lang="en">
<head><meta charset="utf-8"><title>HX-200 FAQ</title>
<style>body { font-family: sans-serif; }</style>
<script>window.analytics = { page: "faq" };</script></head>
<body>
<nav><a href="/">Home</a> <a href="/products">Products</a> <a href="/support">Support</a></nav>
<main>
<h1>HX-200 frequently asked questions</h1>
<h2>Can I verify the calibrator myself?</h2>
<p>Yes. Verification against a reference multimeter is allowed and described in chapter 3 of the
maintenance manual. Adjustment is not: it must be done by an accredited laboratory, otherwise the
calibration certificate is no longer valid.</p>
<h2>How long does the warm-up take?</h2>
<p>Forty-five minutes. Readings taken earlier can drift by several microvolts on the 100 mV range.</p>
<h2>Which firmware version do I have?</h2>
<p>Open Settings, then About. Versions before 3.2.0 do not report the warning W-212 and should be
updated with the tool on the support page.</p>
<img src="status-screen.png" alt="Status screen with the About menu open">
<h2>What does error E-415 mean?</h2>
<p>The contact resistance of the output relay is above its limit. Order relay RLY-4471-B and follow
section 4 of the manual, or book a service visit.</p>
<h2>Where do I find my serial number?</h2>
<p>On the label under the instrument and in Settings, About. It starts with HX2 followed by seven
digits, for example HX20481736.</p>
<img src="serial-label.png" alt="Serial number label under the instrument">
</main>
<footer>Halden Instruments. All product names are fictional.</footer>
</body></html>
"""

CODE = '''"""Nightly job that copies closed tickets into the reporting warehouse."""

from __future__ import annotations

import datetime as dt
from dataclasses import dataclass


@dataclass
class Ticket:
    id: str
    product: str
    priority: str
    opened: dt.datetime
    closed: dt.datetime | None

    @property
    def resolution_hours(self) -> float | None:
        if self.closed is None:
            return None
        return (self.closed - self.opened).total_seconds() / 3600


def first_response_target(priority: str) -> dt.timedelta:
    """Targets from the onboarding handbook, in calendar time."""
    targets = {"P1": dt.timedelta(hours=4), "P2": dt.timedelta(days=1), "P3": dt.timedelta(days=3)}
    return targets[priority]


def breached(ticket: Ticket, first_response: dt.datetime) -> bool:
    return first_response - ticket.opened > first_response_target(ticket.priority)


def export(tickets: list[Ticket]) -> list[dict]:
    rows = []
    for ticket in tickets:
        if ticket.closed is None:
            continue
        rows.append(
            {
                "id": ticket.id,
                "product": ticket.product,
                "priority": ticket.priority,
                "resolution_hours": round(ticket.resolution_hours or 0, 1),
            }
        )
    return rows
'''


def story_pdf(html: str, path: Path, archive_dir: Path) -> None:
    story = pymupdf.Story(html=html, user_css=CSS, archive=pymupdf.Archive(str(archive_dir)))
    writer = pymupdf.DocumentWriter(str(path))
    mediabox = pymupdf.paper_rect("a4")
    where = mediabox + (56, 56, -56, -56)
    more = True
    while more:
        device = writer.begin_page(mediabox)
        more, _ = story.place(where)
        story.draw(device)
        writer.end_page()
    writer.close()
    document = pymupdf.open(str(path))
    document.subset_fonts()
    document.set_metadata({"title": path.stem, "author": "Halden Instruments", "producer": "", "creator": ""})
    document.ez_save(str(path) + ".tmp")
    document.close()
    Path(str(path) + ".tmp").replace(path)


def wiring_png(path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page(width=420, height=220)
    shape = page.new_shape()
    shape.draw_rect(pymupdf.Rect(20, 60, 140, 160))
    shape.draw_rect(pymupdf.Rect(190, 85, 250, 135))
    shape.draw_line((140, 110), (190, 110))
    shape.draw_line((250, 95), (360, 95))
    shape.draw_line((250, 125), (360, 125))
    shape.draw_circle((370, 95), 8)
    shape.draw_circle((370, 125), 8)
    shape.finish(color=(0.2, 0.2, 0.3), width=2)
    shape.commit()
    page.insert_text((32, 115), "Reference board", fontsize=10)
    page.insert_text((207, 114), "K1", fontsize=12)
    page.insert_text((385, 99), "T1", fontsize=10)
    page.insert_text((385, 129), "T2", fontsize=10)
    page.insert_text((20, 30), "HX-200 output stage", fontsize=13)
    page.get_pixmap(dpi=144).save(str(path))
    doc.close()


def label_png(path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page(width=260, height=110)
    page.draw_rect(pymupdf.Rect(8, 8, 252, 102), color=(0.1, 0.1, 0.1), width=1.5)
    page.insert_text((20, 34), "HALDEN INSTRUMENTS", fontsize=11, fontname="hebo")
    page.insert_text((20, 58), "Model HX-200", fontsize=10)
    page.insert_text((20, 84), "S/N HX20481736", fontsize=14, fontname="cour")
    page.get_pixmap(dpi=144).save(str(path))
    doc.close()


def scanned_pdf(path: Path, tmp: Path) -> None:
    source = tmp / "invoice-source.pdf"
    story_pdf(INVOICE, source, tmp)
    rendered = pymupdf.open(source)
    out = pymupdf.open()
    for page in rendered:
        pix = page.get_pixmap(dpi=110, colorspace=pymupdf.csGRAY)
        target = out.new_page(width=page.rect.width, height=page.rect.height)
        target.insert_image(target.rect, pixmap=pix)
    out.save(str(path), garbage=4, deflate=True)
    out.close()
    rendered.close()


def policy_docx(path: Path, image: Path) -> None:
    document = docx.Document()
    document.add_heading("Política de devoluciones y garantía", level=1)
    document.add_paragraph(
        "Esta política explica cómo se tramitan las devoluciones y las reparaciones en garantía de "
        "los instrumentos de Halden Instruments vendidos en España. Se aplica a los pedidos realizados "
        "a partir del 1 de enero de 2025 y sustituye a la versión anterior."
    )
    document.add_heading("Plazo de devolución", level=2)
    document.add_paragraph(
        "El cliente puede devolver un instrumento sin usar en un plazo de treinta días desde la entrega. "
        "El instrumento debe llegar en su embalaje original, con todos los accesorios y con el "
        "certificado de calibración. Los gastos de envío de la devolución corren a cargo del cliente, "
        "salvo que el instrumento llegara dañado o no correspondiera al pedido."
    )
    document.add_heading("Garantía", level=2)
    document.add_paragraph(
        "Todos los instrumentos tienen una garantía de veinticuatro meses. La garantía cubre los "
        "defectos de fabricación y de materiales. No cubre el desgaste normal de piezas como el filtro "
        "de aire o la batería de respaldo, ni los daños causados por un uso fuera de las condiciones "
        "indicadas en el manual."
    )
    for item in (
        "Abra una incidencia en el portal de soporte indicando el número de serie.",
        "Adjunte una foto de la etiqueta y una descripción del fallo.",
        "Espere la autorización de devolución antes de enviar el instrumento.",
        "Envíe el instrumento con el número de autorización visible en la caja.",
    ):
        document.add_paragraph(item, style="List Number")
    document.add_heading("Plazos y costes", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, ("Caso", "Plazo de respuesta", "Coste"), strict=True):
        cell.text = text
    for row in (
        ("Devolución sin usar", "5 días laborables", "Envío a cargo del cliente"),
        ("Reparación en garantía", "10 días laborables", "Sin coste"),
        ("Reparación fuera de garantía", "15 días laborables", "Presupuesto previo"),
        ("Verificación anual", "7 días laborables", "Según tarifa vigente"),
    ):
        cells = table.add_row().cells
        for cell, text in zip(cells, row, strict=True):
            cell.text = text
    document.add_heading("Etiqueta del número de serie", level=2)
    document.add_paragraph(
        "El número de serie empieza por HX2 seguido de siete cifras y aparece en la etiqueta situada "
        "bajo el instrumento, como muestra la imagen."
    )
    document.add_picture(str(image), width=Cm(9))
    document.add_paragraph(
        "Si la etiqueta no es legible, el número de serie también aparece en el menú Ajustes, Acerca de."
    )
    document.core_properties.author = "Halden Instruments"
    document.save(str(path))


PRODUCTS = [
    ("HX-200", "Bench calibrator", "Instruments", 6850.00),
    ("HX-FAN-118", "Fan assembly HX-200", "Spare parts", 84.00),
    ("HX-FLT-021", "Air filter, pack of 5", "Spare parts", 19.50),
    ("RLY-4471-B", "Output relay", "Spare parts", 132.00),
    ("BAT-CR2032-HX", "Backup battery", "Spare parts", 6.90),
    ("CAL-KIT-07", "Calibration kit", "Accessories", 1240.00),
    ("CAL-SVC-12", "Annual verification service", "Services", 420.00),
    ("PRB-110-K", "Test leads, 1 kV", "Accessories", 58.00),
]


def catalog_csv(path: Path) -> None:
    rows = []
    for i in range(40):
        sku, name, category, price = PRODUCTS[i % len(PRODUCTS)]
        variant = i // len(PRODUCTS)
        rows.append(
            [
                f"{sku}-V{variant}" if variant else sku,
                name if not variant else f"{name}, variant {variant}",
                category,
                f"{price + variant * 3.5:.2f}",
                f"2025-{(i % 12) + 1:02d}-{(i % 27) + 1:02d}",
                str((i * 37) % 120),
            ]
        )
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["sku", "name", "category", "price_chf", "last_updated", "stock"])
        writer.writerows(rows)


def stock_xlsx(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Zurich"
    sheet.append(["sku", "bin", "on_hand", "reserved", "reorder_point"])
    for i, (sku, *_rest) in enumerate(PRODUCTS):
        sheet.append([sku, f"ZH-{i + 1:02d}-{(i * 3) % 10}", 10 + i * 7, i % 4, 8])
    second = workbook.create_sheet("Madrid")
    second.append(["sku", "bin", "on_hand", "reserved", "reorder_point"])
    for i, (sku, *_rest) in enumerate(PRODUCTS[:6]):
        second.append([sku, f"MAD-{i + 1:02d}", 4 + i * 3, i % 2, 5])
    workbook.save(str(path))


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp_name:
        tmp = Path(tmp_name)
        wiring_png(OUT / "wiring-diagram.png")
        wiring_png(tmp / "wiring-diagram.png")
        story_pdf(MANUAL, OUT / "hx200-maintenance-manual.pdf", tmp)
        scanned_pdf(OUT / "scanned-invoice.pdf", tmp)
        label_png(tmp / "serial-label.png")
        policy_docx(OUT / "politica-devoluciones.docx", tmp / "serial-label.png")
    catalog_csv(OUT / "product-catalog.csv")
    stock_xlsx(OUT / "warehouse-stock.xlsx")
    (OUT / "onboarding-handbook.md").write_text(HANDBOOK.lstrip(), encoding="utf-8")
    (OUT / "hx200-faq.html").write_text(FAQ, encoding="utf-8")
    (OUT / "ticket_export.py").write_text(CODE, encoding="utf-8")
    print(f"corpus written to {OUT}")


if __name__ == "__main__":
    main()
