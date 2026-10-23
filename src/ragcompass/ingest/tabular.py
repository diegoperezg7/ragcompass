"""CSV and Excel workbooks."""

from __future__ import annotations

import contextlib
import csv
import io
from pathlib import Path

from ragcompass.ingest.text import table_block
from ragcompass.models import Block, BlockType, DocKind


def _clean(rows: list[list[str]]) -> list[list[str]]:
    rows = [[cell.strip() for cell in row] for row in rows]
    rows = [row for row in rows if any(row)]
    width = max((len(row) for row in rows), default=0)
    while width and all(len(row) < width or not row[width - 1] for row in rows):
        width -= 1
    return [row[:width] + [""] * (width - len(row)) for row in rows]


def extract_csv(path: Path):
    from ragcompass.ingest import Extracted, read_text

    source = read_text(path)
    delimiter = "\t" if path.suffix.lower() == ".tsv" else ","
    if delimiter == ",":
        with contextlib.suppress(csv.Error):
            delimiter = csv.Sniffer().sniff(source[:4096], delimiters=",;\t|").delimiter
    rows = _clean(list(csv.reader(io.StringIO(source), delimiter=delimiter)))
    blocks = [table_block(rows)] if rows else []
    return Extracted(kind=DocKind.CSV, blocks=blocks)


def extract_xlsx(path: Path):
    from openpyxl import load_workbook

    from ragcompass.ingest import Extracted

    workbook = load_workbook(path, read_only=True, data_only=True)
    blocks: list[Block] = []
    for sheet in workbook.worksheets:
        rows = _clean(
            [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        )
        if not rows:
            continue
        blocks.append(Block(type=BlockType.HEADING, text=sheet.title, level=1))
        blocks.append(table_block(rows))
    workbook.close()
    return Extracted(kind=DocKind.XLSX, blocks=blocks)
